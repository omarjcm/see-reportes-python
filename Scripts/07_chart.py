from openpyxl.chart import PieChart, Reference
from funciones_excel import cargar_archivo, obtener_hoja

ruta = '../Data/test.xlsx'
wb = cargar_archivo( ruta )
hoja = obtener_hoja(wb, 'Seguidores')

# Código para añadir un Chart en Excel
grafico = PieChart()

categorias = Reference(hoja, min_col=2, min_row=2, max_row=hoja.max_row)
datos = Reference(hoja, min_col=3, min_row=2, max_row=hoja.max_row)

grafico.add_data( datos )
grafico.set_categories( categorias )
grafico.title = 'Seguidores'

hoja.add_chart( grafico, 'E2' )

wb.save( ruta )